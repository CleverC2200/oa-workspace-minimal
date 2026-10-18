import io
import hashlib
import random
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
import openpyxl
import fitz
import pytest
from apps.oa_approval import attachment_extractors


def book(tmp_path: Path, *, broken=False, image_content=None):
    path = tmp_path / 'photos.xlsx'
    image = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 8, 8), False)
    image.clear_with(255)
    png = io.BytesIO(image_content or image.tobytes('png'))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '现场'
    ws['A1'] = '项目'
    ws['B2'] = '=DISPIMG("ID_TEST",1)'
    wb.save(path)
    with zipfile.ZipFile(path) as z:
        parts = {name: z.read(name) for name in z.namelist()}
    parts['xl/worksheets/sheet1.xml'] = parts['xl/worksheets/sheet1.xml'].replace(b'</worksheet>', b'<drawing xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rId1"/></worksheet>')
    with zipfile.ZipFile(path, 'w') as z:
        for name, content in parts.items():
            z.writestr(name, content)
        z.writestr('xl/worksheets/_rels/sheet1.xml.rels', '<Relationships><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing" Target="../drawings/drawing1.xml"/></Relationships>')
        z.writestr('xl/drawings/drawing1.xml', '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><xdr:oneCellAnchor><xdr:from><xdr:col>2</xdr:col><xdr:row>2</xdr:row></xdr:from><xdr:pic><a:blip r:embed="rId1"/></xdr:pic></xdr:oneCellAnchor></xdr:wsDr>')
        z.writestr('xl/drawings/_rels/drawing1.xml.rels', '<Relationships><Relationship Id="rId1" Target="../media/floating.png"/></Relationships>')
        z.writestr('xl/media/floating.png', png.getvalue())
        z.writestr('xl/cellimages.xml', '<root xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><cellImage><xdr:pic><xdr:nvPicPr><xdr:cNvPr name="ID_TEST"/></xdr:nvPicPr><xdr:blipFill><a:blip r:embed="rId1"/></xdr:blipFill></xdr:pic></cellImage></root>')
        z.writestr('xl/_rels/cellimages.xml.rels', '<Relationships><Relationship Id="rId1" Target="media/cell.png"/></Relationships>')
        if not broken:
            z.writestr('xl/media/cell.png', png.getvalue())
    return path


def test_preparation_keeps_wps_cell_and_floating_images(tmp_path):
    path = book(tmp_path)
    images, manifest, facts = attachment_extractors.prepare_materials([{'resource_id':'book', 'path':str(path)}], out_dir=tmp_path/'prepared')
    assert len(images) == 2
    assert {image['spreadsheet_cell'] for image in images} == {'现场!B2', '现场!C3'}
    assert all(image['resource_id']=='book' and Path(image['path']).is_file() for image in images)
    assert manifest[0]['image_count'] == 2
    assert any('B2' in text for text in facts[0]['key_text'])


def test_broken_image_relation_is_not_reported_as_complete(tmp_path):
    path = book(tmp_path, broken=True)
    with pytest.raises(RuntimeError, match='attachment_content_incomplete'):
        attachment_extractors.prepare_materials([{'resource_id':'book', 'path':str(path)}], out_dir=tmp_path/'prepared')


def rewrite_part(path, member, transform):
    with zipfile.ZipFile(path) as z:
        parts = {name: z.read(name) for name in z.namelist()}
    parts[member] = transform(parts[member])
    with zipfile.ZipFile(path, 'w') as z:
        for name, content in parts.items():
            z.writestr(name, content)


@pytest.mark.parametrize('compact', [False, True])
def test_preparation_ignores_unreferenced_workbook_custom_xml(tmp_path, compact):
    path = book(tmp_path)
    namespace = 'http://schemas.openxmlformats.org/package/2006/relationships'

    def add_metadata(content):
        root = ET.fromstring(content)
        ET.SubElement(root, f'{{{namespace}}}Relationship', {
            'Id': 'metadata',
            'Type': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml',
            'Target': '../customXml/item1.xml',
        })
        return ET.tostring(root)

    rewrite_part(path, 'xl/_rels/workbook.xml.rels', add_metadata)
    with zipfile.ZipFile(path, 'a') as archive:
        archive.writestr('customXml/item1.xml', '<metadata/>')
    original = path.read_bytes()

    images, manifest, facts = attachment_extractors.prepare_materials(
        [{'resource_id': 'book', 'path': str(path)}],
        out_dir=tmp_path/'prepared', compact_spreadsheet=compact,
    )

    assert path.read_bytes() == original
    assert {image['spreadsheet_cell'] for image in images} == {'现场!B2', '现场!C3'}
    assert manifest[0]['status'] == 'spreadsheet_read'
    assert manifest[0]['image_count'] == 2
    assert facts[0]['resource_id'] == 'book'


def test_shared_formula_preserves_each_cell_anchor(tmp_path):
    path = book(tmp_path)
    rewrite_part(path, 'xl/worksheets/sheet1.xml', lambda b: b.replace(b'<f>DISPIMG', b'<f t="shared" si="0" ref="B2:C2">DISPIMG').replace(b'</row></sheetData>', b'<c r="C2"><f t="shared" si="0"/></c></row></sheetData>'))
    images, _, _ = attachment_extractors.prepare_materials([{'resource_id':'book', 'path':str(path)}], out_dir=tmp_path/'prepared')
    assert {image['spreadsheet_cell'] for image in images} == {'现场!B2', '现场!C2', '现场!C3'}


def test_corrupt_image_is_not_successfully_prepared(tmp_path):
    path = book(tmp_path)
    rewrite_part(path, 'xl/media/cell.png', lambda _: b'not a PNG')
    with pytest.raises(RuntimeError, match='attachment_content_incomplete'):
        attachment_extractors.prepare_materials([{'resource_id':'book', 'path':str(path)}], out_dir=tmp_path/'prepared')


def test_image_relationship_cannot_escape_workbook(tmp_path):
    path = book(tmp_path)
    rewrite_part(path, 'xl/_rels/cellimages.xml.rels', lambda b: b.replace(b'media/cell.png', b'../../outside.png'))
    with pytest.raises(RuntimeError, match='attachment_content_incomplete'):
        attachment_extractors.prepare_materials([{'resource_id':'book', 'path':str(path)}], out_dir=tmp_path/'prepared')


@pytest.mark.parametrize('member', [
    'xl/_rels/workbook.xml.rels',
    'xl/worksheets/_rels/sheet1.xml.rels',
    'xl/drawings/_rels/drawing1.xml.rels',
])
@pytest.mark.parametrize('mutation', ['external', 'outside', 'missing', 'duplicate'])
def test_preparation_rejects_untrusted_referenced_relationship(tmp_path, member, mutation):
    path = book(tmp_path)

    def mutate(content):
        root = ET.fromstring(content)
        relation = root[0]
        if mutation == 'external':
            relation.set('TargetMode', 'External')
        elif mutation == 'outside':
            relation.set('Target', '../../outside.png')
        elif mutation == 'missing':
            relation.set('Target', 'missing.xml')
        else:
            root.append(ET.fromstring(ET.tostring(relation)))
        return ET.tostring(root)

    rewrite_part(path, member, mutate)
    with pytest.raises(RuntimeError, match='^attachment_content_incomplete$'):
        attachment_extractors.prepare_materials(
            [{'resource_id': 'book', 'path': str(path)}], out_dir=tmp_path/'prepared',
        )


def test_compact_spreadsheet_keeps_pixels_anchors_and_original_source(tmp_path):
    image = fitz.Pixmap(fitz.csRGB, 600, 840, random.Random(7).randbytes(600 * 840 * 3), 0)
    image.set_dpi(144, 144)
    content = image.tobytes('png')
    source = book(tmp_path, image_content=content)
    original = source.read_bytes()
    descriptor = {'resource_id': 'book', 'path': str(source)}
    normal, _, _ = attachment_extractors.prepare_materials([descriptor], out_dir=tmp_path/'normal')
    compact, manifest, facts = attachment_extractors.prepare_materials(
        [descriptor], out_dir=tmp_path/'compact', compact_spreadsheet=True,
    )
    assert source.read_bytes() == original
    assert len(normal) == len(compact) == manifest[0]['image_count'] == 2
    assert {item['spreadsheet_cell'] for item in compact} == {'现场!B2', '现场!C3'}
    assert sum(Path(item['path']).stat().st_size for item in compact) < len(content) * 2 * 0.7
    for before, after in zip(normal, compact):
        assert Path(before['path']).read_bytes() == content
        assert after['resource_id'] == before['resource_id'] == 'book'
        assert after['source_content_sha256'] == hashlib.sha256(original).hexdigest()
        prepared = Path(after['path'])
        assert prepared.suffix == '.jpg'
        assert after['content_sha256'] == hashlib.sha256(prepared.read_bytes()).hexdigest()
        decoded = fitz.Pixmap(str(prepared))
        assert (decoded.width, decoded.height) == (600, 840)
    assert compact[0]['content_sha256'] == compact[1]['content_sha256']
    assert any('B2' in text for text in facts[0]['key_text'])
    assert any('C3' in text for text in facts[0]['key_text'])


@pytest.mark.parametrize('transparent', [False, True])
def test_compact_spreadsheet_keeps_small_or_transparent_images_lossless(tmp_path, transparent):
    if transparent:
        image = fitz.Pixmap(fitz.csRGB, 400, 400, random.Random(8).randbytes(400 * 400 * 4), 1)
    else:
        image = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 8, 8), False)
        image.clear_with(255)
    content = image.tobytes('png')
    source = book(tmp_path, image_content=content)
    images, _, _ = attachment_extractors.prepare_materials(
        [{'resource_id': 'book', 'path': str(source)}],
        out_dir=tmp_path/'prepared', compact_spreadsheet=True,
    )
    assert all(Path(item['path']).read_bytes() == content for item in images)
