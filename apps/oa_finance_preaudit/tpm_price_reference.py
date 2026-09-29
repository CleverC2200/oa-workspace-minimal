"""Read and freeze finance's reference prices; material matching belongs to the reviewer."""

from __future__ import annotations

import argparse
from decimal import Decimal, InvalidOperation
import hashlib
from io import BytesIO
import json
from pathlib import Path
import re

from openpyxl import load_workbook

from apps import json_artifact


SUBJECT = "终端建设费-终端形象制作费"
REFERENCE_PATH = Path("reference/tpm-price-reference.json")
SHEET = "参考报价单"
HEADERS = ("材料名称", "材料编码", "材质说明/要求", "单位", "未税指导价（元）", "含税指导价（元，1%）")


def applies(material: dict) -> bool:
    fields = material.get("declared_fields")
    if isinstance(fields, dict):
        return fields.get("budget_account") == SUBJECT
    return isinstance(fields, list) and any(
        isinstance(field, dict)
        and (field.get("code") == "yskm" or field.get("name") == "预算科目")
        and field.get("read_status") not in ("empty", "unresolved")
        and str(field.get("value") or "").strip() == SUBJECT
        for field in fields
    )


def validate(catalog: object) -> dict:
    if not isinstance(catalog, dict) or (
        catalog.get("schema_version") != 1
        or catalog.get("sheet") != SHEET
        or catalog.get("price_column") != "H"
        or not re.fullmatch(r"[0-9a-f]{64}", str(catalog.get("source_sha256") or ""))
        or not isinstance(catalog.get("rows"), list)
        or not catalog["rows"]
    ):
        raise ValueError("tpm_price_reference_invalid")
    codes: set[str] = set()
    source_rows: set[int] = set()
    for row in catalog["rows"]:
        if not isinstance(row, dict) or set(row) != {
            "source_row", "material_code", "name", "specification", "unit", "guide_price"
        }:
            raise ValueError("tpm_price_reference_invalid")
        if (
            type(row["source_row"]) is not int or row["source_row"] < 2
            or row["source_row"] in source_rows
            or any(not isinstance(row[key], str) or not row[key].strip()
                   for key in ("material_code", "name", "specification", "unit", "guide_price"))
            or row["material_code"] in codes
        ):
            raise ValueError("tpm_price_reference_invalid")
        try:
            price = Decimal(row["guide_price"])
        except InvalidOperation:
            raise ValueError("tpm_price_reference_invalid") from None
        if not price.is_finite() or price <= 0:
            raise ValueError("tpm_price_reference_invalid")
        codes.add(row["material_code"])
        source_rows.add(row["source_row"])
    return catalog


def read_workbook(path: Path) -> dict:
    raw = Path(path).read_bytes()
    workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    try:
        sheet = workbook[SHEET]
        if tuple(sheet.cell(1, col).value for col in range(3, 9)) != HEADERS:
            raise ValueError("tpm_price_reference_headers_invalid")
        rows = []
        for number, values in enumerate(sheet.iter_rows(min_row=2, max_col=8, values_only=True), 2):
            if all(value is None for value in values):
                continue
            name, code, specification, unit, _untaxed, price = values[2:]
            rows.append({
                "source_row": number,
                "material_code": "" if code is None else str(code),
                "name": "" if name is None else str(name),
                "specification": "" if specification is None else str(specification),
                "unit": "" if unit is None else str(unit),
                "guide_price": "" if price is None else str(price),
            })
        return validate({
            "schema_version": 1,
            "source_sha256": hashlib.sha256(raw).hexdigest(),
            "sheet": SHEET,
            "price_column": "H",
            "rows": rows,
        })
    finally:
        workbook.close()


def snapshot(path: Path) -> dict:
    try:
        raw = path.read_bytes()
    except FileNotFoundError:
        return {"status": "not_configured"}
    except OSError:
        return {"status": "unreadable"}
    try:
        catalog = validate(json.loads(raw))
    except (ValueError, TypeError):
        return {"status": "unreadable"}
    return {**catalog, "status": "ready", "catalog_sha256": hashlib.sha256(raw).hexdigest()}


def evidence_gap(item: dict, material: dict) -> str | None:
    if item.get("item_key") != "attachment_completeness" or item.get("status") != "pass" or not applies(material):
        return None
    catalog = material.get("tpm_price_reference") or {}
    if catalog.get("status") != "ready":
        return "参考报价未加载成功，制作清单单价未完成核对，请财务核实。"
    valid_refs = {f"material:tpm_price_reference.rows.{index}" for index in range(len(catalog["rows"]))}
    if not valid_refs.intersection(item.get("evidence_refs") or []):
        return "未引用参考报价的具体物料，制作清单单价未完成核对，请财务核实。"
    return None


def main() -> None:
    parser = argparse.ArgumentParser(description="Import the approved reference sheet into local runtime data.")
    parser.add_argument("source", type=Path)
    parser.add_argument("--runtime-root", type=Path, default=Path("runtime"))
    args = parser.parse_args()
    catalog = read_workbook(args.source)
    destination = args.runtime_root / REFERENCE_PATH
    destination.parent.mkdir(parents=True, exist_ok=True)
    json_artifact.write_json_atomic(destination, catalog, indent=2)
    print(json.dumps({"status": "imported", "row_count": len(catalog["rows"]), "path": str(destination)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
